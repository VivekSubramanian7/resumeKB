"""Smoke test for doc-ingest conversion."""

from __future__ import annotations

import tempfile
from pathlib import Path

import docx

from doc_ingest import UnsupportedDocumentType, convert_to_markdown


def main() -> None:
    d = docx.Document()
    d.add_heading("Vivek Subramanian", level=1)
    d.add_paragraph("Skills: Python, Kubernetes")
    p = Path(tempfile.mkdtemp()) / "cv.docx"
    d.save(str(p))
    md = convert_to_markdown(p)
    assert "Vivek" in md and "Kubernetes" in md
    err: UnsupportedDocumentType | None = None
    try:
        convert_to_markdown(p.with_suffix(".doc"))
    except UnsupportedDocumentType as exc:
        err = exc
    assert err is not None
    print("docingest-ok")


if __name__ == "__main__":
    main()
